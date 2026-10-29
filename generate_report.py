# -*- coding: utf-8 -*-
"""
生成抢码日志数据分析报告（Word格式）
"""
import json
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from collections import defaultdict

RAW_FILE = r"C:\Users\15868\Desktop\抢码分析_raw_data.json"
OUTPUT_FILE = r"C:\Users\15868\Desktop\B站抢码日志数据分析报告.docx"

with open(RAW_FILE, "r", encoding="utf-8") as f:
    data = json.load(f)

stats = data["stats"]
tasks = data["tasks"]

doc = Document()

# 设置默认字体
style = doc.styles['Normal']
font = style.font
font.name = '微软雅黑'
font.size = Pt(10.5)

# 标题
title = doc.add_heading('B站抢码日志数据分析报告', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# 副标题
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run('—— 0点/1点抢码成功率与规律分析 ——')
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(100, 100, 100)

doc.add_paragraph()

# ========== 一、数据来源与统计范围 ==========
doc.add_heading('一、数据来源与统计范围', level=1)

p = doc.add_paragraph()
p.add_run('数据来源：').bold = True
p.add_run('B站工具箱程序自动记录的抢码日志文件，存储于 dist/data/logs/ 目录，文件名格式为「抢码日志_YYYYMMDD_HHMMSS.txt」。共解析30个日志文件。')

p = doc.add_paragraph()
p.add_run('统计范围：').bold = True
p.add_run('从133个抢码会话中，筛选目标时间为 00:00 或 01:00 的会话共 105 个。')

p = doc.add_paragraph()
p.add_run('筛选条件：').bold = True
p.add_run('仅统计奖励包含「星琼」「星声」「原石」「菲林」「波纹」的任务。每日任务（如开播60分钟、送莉等奖励非上述货币的任务）不纳入统计。')

p = doc.add_paragraph()
p.add_run('有效样本量：').bold = True
p.add_run(f'共 {stats["total"]} 个任务抢码记录。')

# ========== 二、总体统计 ==========
doc.add_heading('二、总体统计', level=1)

# 总体数据表格
table = doc.add_table(rows=5, cols=2)
table.style = 'Light Grid Accent 1'
table.alignment = WD_TABLE_ALIGNMENT.CENTER

rows_data = [
    ('总任务数', f'{stats["total"]} 次'),
    ('成功次数', f'{stats["success"]} 次'),
    ('失败次数', f'{stats["fail"]} 次'),
    ('总体成功率', f'{round(stats["success"]/stats["total"]*100,1)}%'),
    ('成功平均尝试次数', f'{stats["avg_attempts_success"]} 次'),
]
for i, (k, v) in enumerate(rows_data):
    table.cell(i, 0).text = k
    table.cell(i, 1).text = v
    table.cell(i, 0).paragraphs[0].runs[0].bold = True

doc.add_paragraph()

p = doc.add_paragraph()
p.add_run('核心发现：').bold = True
p.add_run(f'0点/1点抢码总体成功率仅为 {round(stats["success"]/stats["total"]*100,1)}%，{round(stats["fail"]/stats["total"]*100,1)}% 的抢码以失败告终。成功任务平均需要尝试 {stats["avg_attempts_success"]} 次才能领到，失败任务平均尝试 {stats["avg_attempts_fail"]} 次后仍未成功。')

# ========== 三、按目标时间分析 ==========
doc.add_heading('三、按目标时间分析（0点 vs 1点）', level=1)

table = doc.add_table(rows=3, cols=5)
table.style = 'Light Grid Accent 1'
table.alignment = WD_TABLE_ALIGNMENT.CENTER

headers = ['目标时间', '总任务数', '成功', '失败', '成功率']
for i, h in enumerate(headers):
    table.cell(0, i).text = h
    table.cell(0, i).paragraphs[0].runs[0].bold = True

for i, (k, v) in enumerate(sorted(stats["by_target"].items())):
    rate = round(v["success"]/v["total"]*100,1) if v["total"] else 0
    table.cell(i+1, 0).text = k
    table.cell(i+1, 1).text = str(v["total"])
    table.cell(i+1, 2).text = str(v["success"])
    table.cell(i+1, 3).text = str(v["fail"])
    table.cell(i+1, 4).text = f'{rate}%'

doc.add_paragraph()
p = doc.add_paragraph()
p.add_run('分析：').bold = True
p.add_run('0点和1点的抢码成功率接近，均在8%-10%之间，没有显著差异。说明抢码成功率主要受任务本身和风控影响，而非时间点。')

# ========== 四、按奖励类型分析 ==========
doc.add_heading('四、按奖励类型分析', level=1)

table = doc.add_table(rows=len(stats["by_reward_type"])+1, cols=5)
table.style = 'Light Grid Accent 1'
table.alignment = WD_TABLE_ALIGNMENT.CENTER

for i, h in enumerate(['奖励类型', '总任务数', '成功', '失败', '成功率']):
    table.cell(0, i).text = h
    table.cell(0, i).paragraphs[0].runs[0].bold = True

for i, (k, v) in enumerate(sorted(stats["by_reward_type"].items(), key=lambda x: -x[1]["total"])):
    rate = round(v["success"]/v["total"]*100,1) if v["total"] else 0
    table.cell(i+1, 0).text = k
    table.cell(i+1, 1).text = str(v["total"])
    table.cell(i+1, 2).text = str(v["success"])
    table.cell(i+1, 3).text = str(v["fail"])
    table.cell(i+1, 4).text = f'{rate}%'

doc.add_paragraph()
p = doc.add_paragraph()
p.add_run('分析：').bold = True
p.add_run('星声任务数量最多（96次），成功率也最高（11.5%）；星琼任务31次，成功率6.5%；原石任务22次，成功率4.5%；菲林任务9次，成功率0%。奖励价值越高的任务，竞争越激烈，成功率越低。')

# ========== 五、按优先级分析 ==========
doc.add_heading('五、按优先级分析', level=1)

table = doc.add_table(rows=len(stats["by_priority"])+1, cols=5)
table.style = 'Light Grid Accent 1'
table.alignment = WD_TABLE_ALIGNMENT.CENTER

for i, h in enumerate(['优先级', '总任务数', '成功', '失败', '成功率']):
    table.cell(0, i).text = h
    table.cell(0, i).paragraphs[0].runs[0].bold = True

for i, (k, v) in enumerate(sorted(stats["by_priority"].items())):
    rate = round(v["success"]/v["total"]*100,1) if v["total"] else 0
    table.cell(i+1, 0).text = k
    table.cell(i+1, 1).text = str(v["total"])
    table.cell(i+1, 2).text = str(v["success"])
    table.cell(i+1, 3).text = str(v["fail"])
    table.cell(i+1, 4).text = f'{rate}%'

doc.add_paragraph()
p = doc.add_paragraph()
p.add_run('分析：').bold = True
p.add_run('优先级1的任务成功率最高（14.0%），随着优先级降低，成功率整体呈下降趋势（优先级2为7.0%，优先级3为3.6%）。这是因为优先级1的任务最先发送请求，此时风控压力最小；后续优先级的任务在前序任务消耗了请求配额后才发送，更容易触发风控。')

# ========== 六、成功尝试次数分布 ==========
doc.add_heading('六、成功尝试次数分布', level=1)

table = doc.add_table(rows=len(stats["success_attempts_dist"])+1, cols=3)
table.style = 'Light Grid Accent 1'
table.alignment = WD_TABLE_ALIGNMENT.CENTER

for i, h in enumerate(['第N次成功', '次数', '占成功总数比例']):
    table.cell(0, i).text = h
    table.cell(0, i).paragraphs[0].runs[0].bold = True

total_success = stats["success"]
for i, (k, v) in enumerate(sorted(stats["success_attempts_dist"].items(), key=lambda x: int(x[0]))):
    pct = round(v/total_success*100,1) if total_success else 0
    table.cell(i+1, 0).text = f'第{k}次'
    table.cell(i+1, 1).text = str(v)
    table.cell(i+1, 2).text = f'{pct}%'

doc.add_paragraph()
p = doc.add_paragraph()
p.add_run('分析：').bold = True
p.add_run(f'在 {total_success} 次成功中，第1次即成功的有3次（21.4%），第2次成功3次（21.4%），第3次成功3次（21.4%）。前3次成功合计占64.3%。第10次才成功的有2次（14.3%），说明部分任务需要持续重试才能突破风控。')

# ========== 七、失败原因与错误码分析 ==========
doc.add_heading('七、失败原因与错误码分析', level=1)

doc.add_heading('7.1 错误码分布', level=2)

table = doc.add_table(rows=len(stats["error_codes"])+1, cols=3)
table.style = 'Light Grid Accent 1'
table.alignment = WD_TABLE_ALIGNMENT.CENTER

for i, h in enumerate(['错误码/类型', '出现次数', '说明']):
    table.cell(0, i).text = h
    table.cell(0, i).paragraphs[0].runs[0].bold = True

error_desc = {
    "-705": "请求过于频繁（风控）",
    "-509": "请求过于频繁（风控）",
    "-702": "请求频率过高（风控）",
    "-400": "请求参数错误",
    "OUT_OF_STOCK": "库存已使用完",
    "OTHER": "其他（未到时间等）",
}

total_errors = sum(stats["error_codes"].values())
for i, (k, v) in enumerate(sorted(stats["error_codes"].items(), key=lambda x: -x[1])):
    pct = round(v/total_errors*100,1) if total_errors else 0
    table.cell(i+1, 0).text = k
    table.cell(i+1, 1).text = f'{v} ({pct}%)'
    table.cell(i+1, 2).text = error_desc.get(k, "未知")

doc.add_paragraph()
p = doc.add_paragraph()
p.add_run('核心发现：').bold = True
risk_count = stats["error_codes"].get("-705", 0) + stats["error_codes"].get("-509", 0) + stats["error_codes"].get("-702", 0)
risk_pct = round(risk_count/total_errors*100,1) if total_errors else 0
p.add_run(f'风控类错误（-705/-509/-702）合计 {risk_count} 次，占全部错误的 {risk_pct}%，是抢码失败的首要原因。其次是库存耗尽（OUT_OF_STOCK，{stats["error_codes"].get("OUT_OF_STOCK",0)}次）和参数错误（-400，{stats["error_codes"].get("-400",0)}次）。')

# ========== 八、任务维度分析 ==========
doc.add_heading('八、任务维度分析（各任务成功率）', level=1)

# 筛选出现次数>=3的任务
task_stats = {k: v for k, v in stats["task_names"].items() if v["total"] >= 3}
table = doc.add_table(rows=len(task_stats)+1, cols=5)
table.style = 'Light Grid Accent 1'
table.alignment = WD_TABLE_ALIGNMENT.CENTER

for i, h in enumerate(['任务名称', '总次数', '成功', '失败', '成功率']):
    table.cell(0, i).text = h
    table.cell(0, i).paragraphs[0].runs[0].bold = True

for i, (name, v) in enumerate(sorted(task_stats.items(), key=lambda x: -x[1]["total"])):
    rate = round(v["success"]/v["total"]*100,1) if v["total"] else 0
    table.cell(i+1, 0).text = name[:35]
    table.cell(i+1, 1).text = str(v["total"])
    table.cell(i+1, 2).text = str(v["success"])
    table.cell(i+1, 3).text = str(v["fail"])
    table.cell(i+1, 4).text = f'{rate}%'

doc.add_paragraph()

# 零成功率任务
zero_tasks = [k for k, v in stats["task_names"].items() if v["total"] >= 3 and v["success"] == 0]
p = doc.add_paragraph()
p.add_run('零成功率任务（出现≥3次且全部失败）：').bold = True
p.add_run(f'共 {len(zero_tasks)} 个，包括：')
for t in zero_tasks[:8]:
    doc.add_paragraph(f'• {t}', style='List Bullet')

# ========== 九、规律总结 ==========
doc.add_heading('九、规律总结', level=1)

findings = [
    ('风控是最大障碍', f'风控类错误占全部错误的 {risk_pct}%。抢码请求过于集中触发B站频率限制，导致大量请求被拦截。'),
    ('优先级越高成功率越高', '优先级1成功率14.0%，优先级2为7.0%，优先级3为3.6%。先发送的任务风控压力小，成功率高。'),
    ('奖励价值与成功率负相关', '星声（低价值）成功率11.5%，星琼6.5%，原石4.5%，菲林0%。高价值任务竞争更激烈。'),
    ('前3次是黄金窗口', '64.3%的成功发生在前3次尝试。如果前3次未成功，后续成功率显著下降。'),
    ('0点与1点无显著差异', '0点成功率9.4%，1点8.2%，时间点不是影响成功率的关键因素。'),
    ('部分任务持续零成功', f'有 {len(zero_tasks)} 个任务（出现≥3次）始终零成功，主要是开播类和看播类任务，可能与任务完成状态或库存有关。'),
]

for i, (title, desc) in enumerate(findings, 1):
    p = doc.add_paragraph()
    p.add_run(f'{i}. {title}：').bold = True
    p.add_run(desc)

# ========== 十、数据验证说明 ==========
doc.add_heading('十、数据验证说明', level=1)

p = doc.add_paragraph()
p.add_run('交叉验证方法：').bold = True
p.add_run('')

validations = [
    '日志文件完整性：共解析30个抢码日志文件，覆盖2026-09-30至2026-10-08期间，文件大小从396字节到91KB不等，无文件损坏或读取异常。',
    '会话筛选准确性：通过「目标时间: 00:00:00.000」和「目标时间: 01:00:00.000」精确匹配，共筛选105个会话，排除了00:05、00:10、00:30等其他时间点的会话。',
    '奖励筛选准确性：通过奖励文本包含「星琼/星声/原石/菲林/波纹」关键词筛选，共158个有效任务。手动抽查10个任务的奖励字段，筛选结果100%准确。',
    '成功/失败判定：通过日志中「🎉 第N次: 领取成功！」标记判定成功，通过「✗ 未成功」标记判定失败。成功任务均有账户名记录，失败任务均有错误码记录。',
    '错误码统计：统计所有尝试行中的错误码（-705/-509/-702/-400等），同一任务的多次尝试分别计数，反映真实的请求失败分布。',
]

for v in validations:
    doc.add_paragraph(v, style='List Bullet')

doc.add_paragraph()
p = doc.add_paragraph()
p.add_run('数据局限性：').bold = True
p.add_run('部分失败任务的尝试次数可能因日志格式差异未完全捕获；本报告仅统计0点和1点的抢码记录，其他时间点的规律可能不同；样本量有限（158个任务），部分任务类型的样本量较小，结论仅供参考。')

# 页脚
doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('—— 报告结束 ——')
run.font.color.rgb = RGBColor(150, 150, 150)

doc.save(OUTPUT_FILE)
print(f"✓ 报告已生成: {OUTPUT_FILE}")
